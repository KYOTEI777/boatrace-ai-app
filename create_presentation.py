from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
import pptx.oxml.ns as nsmap
from lxml import etree

# Color palette
NAVY = RGBColor(0x0D, 0x1B, 0x3E)
BLUE = RGBColor(0x1A, 0x4F, 0x9C)
CYAN = RGBColor(0x00, 0xB4, 0xD8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xF0, 0xF4, 0xF8)
GRAY = RGBColor(0x64, 0x74, 0x8B)
ORANGE = RGBColor(0xFF, 0x6B, 0x35)
GREEN = RGBColor(0x06, 0xD6, 0xA0)

prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)

def add_rect(slide, left, top, width, height, fill_color, transparency=0):
    shape = slide.shapes.add_shape(1, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.line.fill.background()
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    return shape

def add_text_box(slide, text, left, top, width, height, font_size=18, bold=False, color=WHITE, align=PP_ALIGN.LEFT, wrap=True):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    txBox.word_wrap = wrap
    tf = txBox.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Meiryo UI"
    return txBox

def add_multiline_text(slide, lines, left, top, width, height, font_size=16, color=WHITE, line_spacing=1.15):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    txBox.word_wrap = True
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, (text, bold, size, clr) in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.space_before = Pt(2)
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size or font_size)
        run.font.bold = bold
        run.font.color.rgb = clr or color
        run.font.name = "Meiryo UI"
    return txBox

# ============================================================
# Slide 1: Title
# ============================================================
slide_layout = prs.slide_layouts[6]  # blank
slide = prs.slides.add_slide(slide_layout)

# Background
add_rect(slide, 0, 0, 13.33, 7.5, NAVY)
# Accent bar
add_rect(slide, 0, 5.8, 13.33, 0.08, CYAN)

# Main title
add_text_box(slide, "ボートレースAI予想アプリ", 1.0, 1.5, 11.33, 1.2,
             font_size=44, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
# Subtitle
add_text_box(slide, "販売戦略プレゼンテーション", 1.0, 2.9, 11.33, 0.8,
             font_size=28, bold=False, color=CYAN, align=PP_ALIGN.CENTER)
# Tagline
add_text_box(slide, "迷わせない予想AI。全場対応、1点勝負。", 1.0, 4.0, 11.33, 0.7,
             font_size=20, bold=False, color=LIGHT_GRAY, align=PP_ALIGN.CENTER)
# Date
add_text_box(slide, "2026年3月", 10.5, 6.8, 2.5, 0.5,
             font_size=14, color=GRAY, align=PP_ALIGN.RIGHT)

# ============================================================
# Slide 2: Agenda
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(slide, 0, 0, 13.33, 7.5, LIGHT_GRAY)
add_rect(slide, 0, 0, 13.33, 1.3, NAVY)
add_rect(slide, 0, 1.3, 0.08, 6.2, CYAN)

add_text_box(slide, "目次", 0.5, 0.2, 12, 0.9, font_size=34, bold=True, color=WHITE)

agenda_items = [
    "01  サービス概要・強み",
    "02  ターゲットユーザー",
    "03  差別化ポイント",
    "04  価格・プラン設計",
    "05  集客チャネル戦略",
    "06  KPI・重要指標",
    "07  アクションプラン",
]
for i, item in enumerate(agenda_items):
    y = 1.6 + i * 0.7
    add_rect(slide, 0.5, y, 12.0, 0.55, WHITE)
    add_text_box(slide, item, 0.7, y + 0.05, 11.5, 0.5, font_size=18, bold=(i % 2 == 0), color=NAVY)

# ============================================================
# Slide 3: Service Overview
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(slide, 0, 0, 13.33, 7.5, NAVY)
add_rect(slide, 0, 0, 13.33, 1.3, BLUE)
add_rect(slide, 0, 1.3, 13.33, 0.06, CYAN)

add_text_box(slide, "01  サービス概要・強み", 0.5, 0.2, 12, 0.9, font_size=30, bold=True, color=WHITE)

features = [
    ("全場対応", "全国24場すべてのレースに対応", CYAN),
    ("1点予想", "迷いをなくす単一予想で潔く勝負", ORANGE),
    ("AI精度", "機械学習による高精度な予想エンジン", GREEN),
    ("本命〜大穴", "オッズ帯に関わらず幅広くカバー", CYAN),
]

for i, (title, desc, color) in enumerate(features):
    col = i % 2
    row = i // 2
    x = 0.5 + col * 6.4
    y = 1.7 + row * 2.5
    add_rect(slide, x, y, 6.0, 2.1, BLUE)
    add_rect(slide, x, y, 6.0, 0.55, color)
    add_text_box(slide, title, x + 0.2, y + 0.08, 5.6, 0.45, font_size=20, bold=True, color=NAVY)
    add_text_box(slide, desc, x + 0.2, y + 0.65, 5.6, 1.3, font_size=16, color=WHITE)

# ============================================================
# Slide 4: Target Users
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(slide, 0, 0, 13.33, 7.5, LIGHT_GRAY)
add_rect(slide, 0, 0, 13.33, 1.3, NAVY)
add_rect(slide, 0, 1.3, 13.33, 0.06, CYAN)

add_text_box(slide, "02  ターゲットユーザー", 0.5, 0.2, 12, 0.9, font_size=30, bold=True, color=WHITE)

# Main target
add_rect(slide, 0.4, 1.6, 5.8, 5.2, NAVY)
add_rect(slide, 0.4, 1.6, 5.8, 0.6, BLUE)
add_text_box(slide, "メインターゲット", 0.6, 1.65, 5.4, 0.5, font_size=18, bold=True, color=CYAN)
add_multiline_text(slide, [
    ("中級者〜ベテラン", True, 22, WHITE),
    ("", False, 8, WHITE),
    ("• 「情報が多すぎて迷う」という悩みを持つ人", False, 15, LIGHT_GRAY),
    ("• 既存サービスの複数予想に不満がある層", False, 15, LIGHT_GRAY),
    ("• 回収率を真剣に考えているユーザー", False, 15, LIGHT_GRAY),
    ("", False, 8, WHITE),
    ("推定人口：競艇利用者の約40%", False, 14, CYAN),
], 0.6, 2.35, 5.4, 4.2)

# Sub target
add_rect(slide, 6.8, 1.6, 5.8, 5.2, NAVY)
add_rect(slide, 6.8, 1.6, 5.8, 0.6, RGBColor(0x2D, 0x6A, 0xB0))
add_text_box(slide, "サブターゲット", 7.0, 1.65, 5.4, 0.5, font_size=18, bold=True, color=ORANGE)
add_multiline_text(slide, [
    ("初心者", True, 22, WHITE),
    ("", False, 8, WHITE),
    ("• 「どれを買えばいいか\n  わからない」という人", False, 15, LIGHT_GRAY),
    ("• 1点に絞られているため\n  入門ハードルが低い", False, 15, LIGHT_GRAY),
    ("• ボートレース人口拡大の\n  取り込み対象", False, 15, LIGHT_GRAY),
    ("", False, 8, WHITE),
    ("推定人口：新規参入者の約30%", False, 14, ORANGE),
], 7.0, 2.35, 5.4, 4.2)

# ============================================================
# Slide 5: Differentiation
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(slide, 0, 0, 13.33, 7.5, NAVY)
add_rect(slide, 0, 0, 13.33, 1.3, BLUE)
add_rect(slide, 0, 1.3, 13.33, 0.06, ORANGE)

add_text_box(slide, "03  差別化ポイント", 0.5, 0.2, 12, 0.9, font_size=30, bold=True, color=WHITE)

# VS comparison
add_rect(slide, 0.4, 1.6, 5.8, 5.4, RGBColor(0x3D, 0x10, 0x10))
add_text_box(slide, "競合サービス", 0.6, 1.7, 5.4, 0.6, font_size=20, bold=True, color=RGBColor(0xFF, 0x80, 0x80), align=PP_ALIGN.CENTER)
add_multiline_text(slide, [
    ("✗  複数点予想で迷わせる", False, 16, RGBColor(0xFF, 0xAA, 0xAA)),
    ("✗  一部の場のみ対応", False, 16, RGBColor(0xFF, 0xAA, 0xAA)),
    ("✗  本命寄りの無難な予想", False, 16, RGBColor(0xFF, 0xAA, 0xAA)),
    ("✗  的中実績が不透明", False, 16, RGBColor(0xFF, 0xAA, 0xAA)),
    ("✗  情報量が多すぎる", False, 16, RGBColor(0xFF, 0xAA, 0xAA)),
], 0.7, 2.5, 5.3, 3.8)

add_text_box(slide, "VS", 5.9, 3.7, 1.5, 0.8, font_size=32, bold=True, color=ORANGE, align=PP_ALIGN.CENTER)

add_rect(slide, 7.1, 1.6, 5.8, 5.4, RGBColor(0x05, 0x2E, 0x1A))
add_text_box(slide, "本アプリ", 7.3, 1.7, 5.4, 0.6, font_size=20, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
add_multiline_text(slide, [
    ("✓  1点予想で意思決定を支援", False, 16, RGBColor(0x80, 0xFF, 0xC0)),
    ("✓  全国24場すべて対応", False, 16, RGBColor(0x80, 0xFF, 0xC0)),
    ("✓  本命〜大穴まで幅広い", False, 16, RGBColor(0x80, 0xFF, 0xC0)),
    ("✓  的中率・回収率を公開", False, 16, RGBColor(0x80, 0xFF, 0xC0)),
    ("✓  シンプルで使いやすいUI", False, 16, RGBColor(0x80, 0xFF, 0xC0)),
], 7.3, 2.5, 5.3, 3.8)

# ============================================================
# Slide 6: Pricing Plans
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(slide, 0, 0, 13.33, 7.5, LIGHT_GRAY)
add_rect(slide, 0, 0, 13.33, 1.3, NAVY)
add_rect(slide, 0, 1.3, 13.33, 0.06, CYAN)

add_text_box(slide, "04  価格・プラン設計", 0.5, 0.2, 12, 0.9, font_size=30, bold=True, color=WHITE)

plans = [
    ("無料プラン", "¥0", "月/永続", GRAY, [
        "1日1レースのみ閲覧",
        "予想精度を体験できる",
        "広告表示あり",
    ], False),
    ("スタンダード", "¥1,980", "月額（税込）", BLUE, [
        "全開催・全レース閲覧",
        "過去予想の履歴閲覧",
        "メール通知機能",
    ], False),
    ("プレミアム", "¥3,980", "月額（税込）", ORANGE, [
        "スタンダード全機能",
        "回収率レポート",
        "的中アラート通知",
        "優先サポート",
    ], True),
]

for i, (name, price, period, color, features, recommended) in enumerate(plans):
    x = 0.5 + i * 4.2
    add_rect(slide, x, 1.6, 3.9, 5.5, NAVY)
    add_rect(slide, x, 1.6, 3.9, 0.65, color)
    if recommended:
        add_rect(slide, x + 2.3, 1.55, 1.5, 0.4, ORANGE)
        add_text_box(slide, "おすすめ", x + 2.35, 1.57, 1.4, 0.35, font_size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text_box(slide, name, x + 0.15, 1.68, 3.6, 0.55, font_size=19, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text_box(slide, price, x + 0.15, 2.4, 3.6, 0.75, font_size=32, bold=True, color=color if color != GRAY else CYAN, align=PP_ALIGN.CENTER)
    add_text_box(slide, period, x + 0.15, 3.15, 3.6, 0.4, font_size=13, color=GRAY, align=PP_ALIGN.CENTER)
    for j, feat in enumerate(features):
        add_text_box(slide, f"• {feat}", x + 0.2, 3.65 + j * 0.6, 3.5, 0.55, font_size=14, color=LIGHT_GRAY)

# Annual discount note
add_rect(slide, 0.5, 7.1, 12.3, 0.3, RGBColor(0xE0, 0xF0, 0xFF))
add_text_box(slide, "※ 年払いプランは2ヶ月分無料（約17%割引）　例：スタンダード年払い ¥19,800/年", 0.7, 7.12, 12, 0.28, font_size=12, color=NAVY)

# ============================================================
# Slide 7: Marketing Channels
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(slide, 0, 0, 13.33, 7.5, NAVY)
add_rect(slide, 0, 0, 13.33, 1.3, BLUE)
add_rect(slide, 0, 1.3, 13.33, 0.06, GREEN)

add_text_box(slide, "05  集客チャネル戦略", 0.5, 0.2, 12, 0.9, font_size=30, bold=True, color=WHITE)

channels = [
    ("X（Twitter）", "優先度 ★★★", CYAN,
     "ボートレースクラスタが最も濃いSNS\n• 毎日「今日の注目レース」を無料発信\n• 的中結果ツイートで信頼構築\n• ハッシュタグ活用で自然な拡散"),
    ("YouTube / TikTok", "優先度 ★★☆", ORANGE,
     "解説動画で初心者層を取り込む\n• レース解説・予想根拠を動画化\n• 的中シーンのショート動画\n• チャンネル登録からアプリ誘導"),
    ("SEO / ブログ", "優先度 ★★☆", GREEN,
     "検索流入で長期的な集客\n• 「ボートレース 予想 AI」で上位表示\n• 「競艇 点数絞る」等のKWを狙う\n• 無料体験への導線設置"),
]

for i, (ch, priority, color, desc) in enumerate(channels):
    x = 0.4 + i * 4.3
    add_rect(slide, x, 1.6, 4.0, 5.4, RGBColor(0x10, 0x2A, 0x5E))
    add_rect(slide, x, 1.6, 4.0, 0.65, color)
    add_text_box(slide, ch, x + 0.15, 1.65, 3.7, 0.55, font_size=18, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    add_text_box(slide, priority, x + 0.15, 2.35, 3.7, 0.4, font_size=14, bold=True, color=color, align=PP_ALIGN.CENTER)
    add_text_box(slide, desc, x + 0.2, 2.85, 3.6, 4.0, font_size=13, color=LIGHT_GRAY)

# ============================================================
# Slide 8: KPIs
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(slide, 0, 0, 13.33, 7.5, LIGHT_GRAY)
add_rect(slide, 0, 0, 13.33, 1.3, NAVY)
add_rect(slide, 0, 1.3, 13.33, 0.06, CYAN)

add_text_box(slide, "06  KPI・重要指標", 0.5, 0.2, 12, 0.9, font_size=30, bold=True, color=WHITE)

kpis = [
    ("無料→有料\n転換率", "5〜10%", "目標値", CYAN),
    ("月次\n解約率", "5%以下", "目標値", GREEN),
    ("月間\nアクティブ率", "70%以上", "目標値", ORANGE),
    ("的中率\n（公開指標）", "公開・透明化", "信頼構築", CYAN),
]

for i, (label, value, sublabel, color) in enumerate(kpis):
    x = 0.5 + i * 3.1
    add_rect(slide, x, 1.7, 2.8, 2.5, NAVY)
    add_rect(slide, x, 1.7, 2.8, 0.6, color)
    add_text_box(slide, label, x + 0.1, 1.73, 2.6, 0.55, font_size=14, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    add_text_box(slide, value, x + 0.1, 2.4, 2.6, 0.9, font_size=22, bold=True, color=color, align=PP_ALIGN.CENTER)
    add_text_box(slide, sublabel, x + 0.1, 3.35, 2.6, 0.4, font_size=12, color=GRAY, align=PP_ALIGN.CENTER)

# Revenue projection
add_rect(slide, 0.4, 4.5, 12.4, 2.6, NAVY)
add_text_box(slide, "収益シミュレーション（スタンダードプランのみ）", 0.6, 4.6, 12, 0.5, font_size=18, bold=True, color=CYAN)

milestones = [
    ("3ヶ月後", "有料100名", "¥198,000/月"),
    ("6ヶ月後", "有料300名", "¥594,000/月"),
    ("12ヶ月後", "有料1,000名", "¥1,980,000/月"),
]
for i, (period, users, revenue) in enumerate(milestones):
    x = 0.8 + i * 4.1
    add_text_box(slide, period, x, 5.2, 3.5, 0.4, font_size=14, color=GRAY, align=PP_ALIGN.CENTER)
    add_text_box(slide, users, x, 5.65, 3.5, 0.45, font_size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text_box(slide, revenue, x, 6.15, 3.5, 0.45, font_size=20, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
    if i < 2:
        add_text_box(slide, "→", x + 3.3, 5.7, 0.6, 0.45, font_size=22, bold=True, color=CYAN, align=PP_ALIGN.CENTER)

# ============================================================
# Slide 9: Action Plan
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(slide, 0, 0, 13.33, 7.5, NAVY)
add_rect(slide, 0, 0, 13.33, 1.3, BLUE)
add_rect(slide, 0, 1.3, 13.33, 0.06, ORANGE)

add_text_box(slide, "07  アクションプラン", 0.5, 0.2, 12, 0.9, font_size=30, bold=True, color=WHITE)

phases = [
    ("Phase 1\n〜1ヶ月", CYAN, [
        "無料プランのリリース",
        "X アカウント開設・毎日投稿開始",
        "的中実績の記録・公開",
        "フィードバック収集",
    ]),
    ("Phase 2\n〜3ヶ月", ORANGE, [
        "有料プラン公開",
        "X フォロワー500名目標",
        "LP（ランディングページ）作成",
        "YouTube チャンネル開設",
    ]),
    ("Phase 3\n〜6ヶ月", GREEN, [
        "SEO記事の本格展開",
        "プレミアムプラン追加",
        "アフィリエイト・紹介制度",
        "有料会員300名突破目標",
    ]),
]

for i, (phase, color, actions) in enumerate(phases):
    x = 0.4 + i * 4.3
    add_rect(slide, x, 1.6, 4.0, 5.5, RGBColor(0x0A, 0x20, 0x50))
    add_rect(slide, x, 1.6, 4.0, 0.85, color)
    add_text_box(slide, phase, x + 0.1, 1.62, 3.8, 0.82, font_size=18, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    for j, action in enumerate(actions):
        add_rect(slide, x + 0.2, 2.6 + j * 1.0, 3.6, 0.7, RGBColor(0x10, 0x2A, 0x5E))
        add_text_box(slide, f"  {action}", x + 0.3, 2.65 + j * 1.0, 3.4, 0.6, font_size=14, color=WHITE)

# ============================================================
# Slide 10: Closing
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_rect(slide, 0, 0, 13.33, 7.5, NAVY)
add_rect(slide, 0, 6.0, 13.33, 1.5, BLUE)

add_text_box(slide, "迷わせない予想AI。", 1.0, 1.5, 11.33, 1.2,
             font_size=48, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
add_text_box(slide, "全場対応、1点勝負。", 1.0, 2.8, 11.33, 1.0,
             font_size=40, bold=True, color=CYAN, align=PP_ALIGN.CENTER)

add_rect(slide, 4.0, 4.1, 5.33, 0.06, ORANGE)

add_text_box(slide, "信頼が最大の購買動機", 1.0, 4.4, 11.33, 0.7,
             font_size=22, color=LIGHT_GRAY, align=PP_ALIGN.CENTER)
add_text_box(slide, "的中実績の公開・透明性・継続投稿から始めましょう", 1.0, 5.1, 11.33, 0.6,
             font_size=18, color=GRAY, align=PP_ALIGN.CENTER)

add_text_box(slide, "ボートレースAI予想アプリ　販売戦略", 1.0, 6.2, 11.33, 0.5,
             font_size=16, color=WHITE, align=PP_ALIGN.CENTER)

# Save
output_path = "/home/user/boatrace-ai-app/販売戦略プレゼン.pptx"
prs.save(output_path)
print(f"Saved: {output_path}")
